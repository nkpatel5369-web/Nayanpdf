"""Synthetic test documents. No real client data is used anywhere in the tests."""

from __future__ import annotations

import hashlib
import io
import os
import random
import shutil
import subprocess
import sys

import pymupdf
import pytest

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from nayanpdf_engine.registry import Context  # noqa: E402
from nayanpdf_engine.server import _load_ops  # noqa: E402
from nayanpdf_engine.util import tools  # noqa: E402

_load_ops()

FIX = os.path.join(os.path.dirname(__file__), "_fixtures")
OUT = os.path.join(os.path.dirname(__file__), "_out")

BANK_ROWS = [
    # date, description, ref, debit, credit, balance
    ("01/04/2025", "Opening balance", "", "", "", "5,00,000.00 Cr"),
    ("03/04/2025", "NEFT-HDFC0001234-ACME SUPPLIES", "000987654321", "1,23,456.78", "", "3,76,543.22 Cr"),
    ("05/04/2025", "UPI/CR/412345678901/Client A", "412345678901", "", "45,000.00", "4,21,543.22 Cr"),
    ("10/04/2025", "GST PAYMENT CHALLAN 27ABCDE1234F1Z5", "CIN0042", "(12,345.00)", "", "4,09,198.22 Cr"),
    ("15/04/2025", "Interest reversal", "INT-REV-77", "", "1,050.50", "4,10,248.72 Cr"),
]


def sha256(path: str) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        h.update(fh.read())
    return h.hexdigest()


def _numbered_pdf(path: str, n: int, label: str, size=(595, 842)):
    d = pymupdf.open()
    for i in range(n):
        p = d.new_page(width=size[0], height=size[1])
        p.insert_text((72, 100), f"{label} page {i + 1}", fontsize=20)
        p.insert_text((72, 140), f"Marker {label}-{i + 1:03d}", fontsize=12)
        p.draw_rect(pymupdf.Rect(72, 160, 300, 200), color=(0.2, 0.4, 0.8), width=1.5)
    d.set_metadata({"title": f"{label} test document", "author": "NayanPDF Test Suite"})
    d.save(path)
    d.close()


def _financial_docx(path: str):
    import docx
    from docx.shared import Pt

    d = docx.Document()
    d.add_heading("ACME Traders Private Limited", 0)
    d.add_paragraph("Statement of Profit and Loss for the year ended 31 March 2026 (amounts in Rs.)")
    t = d.add_table(rows=5, cols=3)
    t.style = "Table Grid"
    rows = [("Particulars", "FY 2025-26", "FY 2024-25"),
            ("Revenue from operations", "1,23,456.78", "98,765.43"),
            ("Other expenses", "(45,000.00)", "(40,500.10)"),
            ("Depreciation", "12,000.00", "11,250.00"),
            ("Net profit", "66,456.78", "47,015.33")]
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            t.cell(i, j).text = v
    d.add_paragraph("Confidential: PAN ABCDE1234F, GSTIN 27ABCDE1234F1Z5, Account 000123456789, contact accounts@acme.example.")
    d.add_page_break()
    d.add_heading("Bank statement extract - Account 000123456789", 1)
    t2 = d.add_table(rows=1, cols=6)
    t2.style = "Table Grid"
    for j, h in enumerate(("Date", "Description", "Ref No", "Debit", "Credit", "Balance")):
        t2.cell(0, j).text = h
    # Enough rows to force the table across a page boundary.
    for k in range(18):
        for r in BANK_ROWS:
            cells = t2.add_row().cells
            for j, v in enumerate(r):
                cells[j].text = v
    for p in d.paragraphs:
        for run in p.runs:
            run.font.size = Pt(11)
    d.save(path)


def _soffice_convert(src: str, outdir: str) -> str:
    exe = tools.soffice_exe()
    profile = os.path.join(outdir, ".lo-profile")
    subprocess.run([exe, f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "pdf",
                    "--outdir", outdir, src], check=True, capture_output=True, timeout=180)
    return os.path.join(outdir, os.path.splitext(os.path.basename(src))[0] + ".pdf")


def _scanned_from(src_pdf: str, path: str, dpi: int = 200, rotate_page: int | None = None):
    src = pymupdf.open(src_pdf)
    d = pymupdf.open()
    for k, p in enumerate(src):
        pix = p.get_pixmap(dpi=dpi, colorspace=pymupdf.csGRAY)
        np_ = d.new_page(width=p.rect.width, height=p.rect.height)
        np_.insert_image(np_.rect, stream=pix.tobytes("png"))
    d.save(path)
    d.close()
    src.close()


def _image_heavy(path: str):
    from PIL import Image, ImageDraw, ImageFilter

    random.seed(7)
    d = pymupdf.open()
    for k in range(3):
        w, h = 2400, 3200
        im = Image.new("RGB", (w, h))
        dr = ImageDraw.Draw(im)
        for y in range(0, h, 4):
            c = int(80 + 120 * y / h)
            dr.rectangle((0, y, w, y + 4), fill=(c, 140 - k * 30, 255 - c))
        for _ in range(250):
            x, y = random.randint(0, w), random.randint(0, h)
            r = random.randint(20, 200)
            dr.ellipse((x - r, y - r, x + r, y + r), fill=tuple(random.randint(0, 255) for _ in range(3)))
        im = im.filter(ImageFilter.GaussianBlur(1.2))
        noise = Image.effect_noise((w, h), 25).convert("RGB")
        im = Image.blend(im, noise, 0.12)
        buf = io.BytesIO()
        im.save(buf, "JPEG", quality=95)
        p = d.new_page()
        p.insert_image(pymupdf.Rect(36, 120, 559, 806), stream=buf.getvalue())
        p.insert_text((36, 80), f"Image heavy page {k + 1} - text must survive compression", fontsize=12)
    d.save(path)
    d.close()


def _form(path: str):
    """AcroForm with text, checkbox, combo box and radio group (built with reportlab, a test-only dependency)."""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(path, pagesize=A4)
    c.setFont("Helvetica", 16)
    c.drawString(72, 780, "Client onboarding form")
    f = c.acroForm
    f.textfield(name="client_name", x=72, y=730, width=230, height=20, value="")
    f.checkbox(name="gst_registered", x=72, y=690, size=18, checked=False)
    f.choice(name="entity_type", x=72, y=650, width=180, height=20, value="Individual",
             options=["Individual", "Company", "LLP", "Partnership"])
    f.radio(name="filing", value="Monthly", selected=False, x=72, y=610, size=18)
    f.radio(name="filing", value="Quarterly", selected=False, x=172, y=610, size=18)
    c.save()


@pytest.fixture(scope="session")
def fx():
    """Build all fixtures once; return a dict of paths and their original hashes."""
    if os.path.isdir(FIX):
        shutil.rmtree(FIX)
    os.makedirs(FIX)
    os.makedirs(OUT, exist_ok=True)
    f = {}
    for name, n in (("a", 2), ("b", 3), ("c", 4)):
        f[name] = os.path.join(FIX, f"{name}.pdf")
        _numbered_pdf(f[name], n, name.upper())
    f["twenty"] = os.path.join(FIX, "twenty.pdf")
    _numbered_pdf(f["twenty"], 20, "Report")
    f["large"] = os.path.join(FIX, "large.pdf")
    _numbered_pdf(f["large"], 300, "Large")
    docx_path = os.path.join(FIX, "financial.docx")
    _financial_docx(docx_path)
    f["financial_docx"] = docx_path
    if tools.soffice_exe():
        f["financial"] = _soffice_convert(docx_path, FIX)
    f["scanned"] = os.path.join(FIX, "scanned.pdf")
    _scanned_from(f.get("financial") or f["twenty"], f["scanned"], dpi=200)
    f["images"] = os.path.join(FIX, "image_heavy.pdf")
    _image_heavy(f["images"])
    f["form"] = os.path.join(FIX, "form.pdf")
    _form(f["form"])
    f["encrypted"] = os.path.join(FIX, "encrypted.pdf")
    d = pymupdf.open(f["a"])
    d.save(f["encrypted"], encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="open123", owner_pw="owner456",
           permissions=pymupdf.PDF_PERM_PRINT)
    d.close()
    f["restricted"] = os.path.join(FIX, "restricted.pdf")
    d = pymupdf.open(f["a"])
    d.save(f["restricted"], encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="", owner_pw="owner456",
           permissions=pymupdf.PDF_PERM_PRINT)
    d.close()
    f["malformed"] = os.path.join(FIX, "malformed.pdf")
    data = open(f["twenty"], "rb").read()
    idx = data.rfind(b"xref")
    broken = data[:idx] + b"xref\n0 999\ngarbage\n" + data[idx + 40:] if idx > 0 else data[: len(data) // 2]
    with open(f["malformed"], "wb") as fh:
        fh.write(broken)
    f["truncated"] = os.path.join(FIX, "truncated.pdf")
    with open(f["truncated"], "wb") as fh:
        fh.write(data[:200])
    f["notpdf"] = os.path.join(FIX, "notes.pdf")
    with open(f["notpdf"], "w") as fh:
        fh.write("This is a text file with a .pdf extension.\n")
    f["png"] = os.path.join(FIX, "logo.png")
    from PIL import Image, ImageDraw

    im = Image.new("RGBA", (400, 200), (255, 255, 255, 0))
    ImageDraw.Draw(im).rectangle((10, 10, 390, 190), outline=(200, 30, 30, 255), width=8)
    im.save(f["png"])
    f["jpg"] = os.path.join(FIX, "photo.jpg")
    Image.new("RGB", (1200, 900), (30, 120, 200)).save(f["jpg"], quality=90)
    f["txt"] = os.path.join(FIX, "notes.txt")
    with open(f["txt"], "w", encoding="utf-8") as fh:
        fh.write("NayanPDF text to PDF test\n" + "\n".join(f"Line {i}: amount ₹ {i * 1000:,}" for i in range(1, 120)))
    hashes = {k: sha256(v) for k, v in f.items() if os.path.isfile(v)}
    f["_hashes"] = hashes
    yield f


@pytest.fixture()
def ctx():
    c = Context("test")
    yield c
    c.cleanup()


def out(name: str) -> str:
    os.makedirs(OUT, exist_ok=True)
    p = os.path.join(OUT, name)
    if os.path.exists(p):
        os.remove(p)
    return p


def page_texts(path: str, password: str | None = None) -> list[str]:
    d = pymupdf.open(path)
    if d.needs_pass:
        d.authenticate(password or "")
    t = [p.get_text("text") for p in d]
    d.close()
    return t


needs_soffice = pytest.mark.skipif(not tools.soffice_exe(), reason="LibreOffice not available")
needs_tesseract = pytest.mark.skipif(not tools.tesseract_exe(), reason="Tesseract not available")
needs_gs = pytest.mark.skipif(not tools.ghostscript_exe(), reason="Ghostscript not available")
