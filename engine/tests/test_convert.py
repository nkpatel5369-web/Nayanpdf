"""Compression, conversion and OCR tests (spec section 30)."""

import os

import pymupdf
import pytest

from conftest import BANK_ROWS, needs_soffice, needs_tesseract, out, page_texts
from nayanpdf_engine.ops import compress, convert, create, ocr
from nayanpdf_engine.util.numbers import parse_date, parse_number


# ---- number safety -----------------------------------------------------------

@pytest.mark.parametrize("text,value", [
    ("1,23,456.78", 123456.78), ("123,456.78", 123456.78), ("(45,000.00)", -45000.0), ("1,234.00-", -1234.0),
    ("-12.5", -12.5), ("₹ 5,000", 5000), ("Rs. 1,000.50", 1000.5), ("0.75", 0.75), ("12.5%", 0.125),
    ("5,00,000.00 Cr", 500000.0),
])
def test_parse_number_converts(text, value):
    v, fmt = parse_number(text)
    assert abs(v - value) < 1e-9


@pytest.mark.parametrize("text", [
    "000123456789", "412345678901", "INV-2026/0042", "2025-26", "27ABCDE1234F1Z5", "98765 43210",
    "12,34,5", "1,2345.00", "1234567890123456.00", "(12", "--5", "ABCDE1234F",
])
def test_parse_number_keeps_identifiers_as_text(text):
    assert parse_number(text) is None


def test_drcr_marker_preserved_in_format():
    v, fmt = parse_number("5,00,000.00 Cr")
    assert v == 500000.0 and '"Cr"' in fmt
    v, fmt = parse_number("1,000.00 Dr")
    assert v == 1000.0 and '"Dr"' in fmt  # sign is NOT inferred


def test_parse_date():
    import datetime as dt

    assert parse_date("15/04/2025")[0] == dt.date(2025, 4, 15)
    assert parse_date("03/04/2025", "MDY")[0] == dt.date(2025, 3, 4)
    assert parse_date("15-Apr-2025")[0] == dt.date(2025, 4, 15)
    assert parse_date("31/02/2025") is None


# ---- compression -------------------------------------------------------------

def test_compress_image_heavy_reports_actual_sizes(fx, ctx):
    r = compress.compress(ctx, fx["images"], out("compressed.pdf"), preset="balanced")
    actual = os.path.getsize(r["output"])
    orig = os.path.getsize(fx["images"])
    assert r["finalSize"] == actual and r["originalSize"] == orig
    assert r["reductionBytes"] == orig - actual
    assert abs(r["reductionPercent"] - round((orig - actual) / orig * 100, 2)) < 0.01
    assert r["finalSize"] < r["originalSize"] * 0.6, r
    t = page_texts(r["output"])
    assert all("text must survive compression" in x for x in t)


def test_compress_presets_and_larger_is_reported(fx, ctx):
    r = compress.compress(ctx, fx["a"], out("tiny_max.pdf"), preset="maximum")
    assert r["larger"] == (r["finalSize"] > r["originalSize"])
    r2 = compress.compress(ctx, fx["images"], out("strong.pdf"), preset="strong")
    r3 = compress.compress(ctx, fx["images"], out("high.pdf"), preset="high")
    assert r2["finalSize"] <= r3["finalSize"]


def test_compress_target_size(fx, ctx):
    r = compress.compress(ctx, fx["images"], out("target.pdf"), targetSizeKb=400)
    assert r["attempts"]
    if r["targetReached"]:
        assert r["finalSize"] <= 400 * 1024


# ---- PDF -> Word ---------------------------------------------------------------

@needs_soffice
def test_pdf_to_word_layout(fx, ctx):
    import docx

    r = convert.to_word(ctx, fx["financial"], out("fin_layout.docx"), mode="layout", pages="1")
    d = docx.Document(r["output"])
    text = "\n".join(p.text for p in d.paragraphs)
    cells = [c.text for t in d.tables for row in t.rows for c in row.cells]
    assert "ACME Traders" in text
    assert any("1,23,456.78" in c for c in cells), cells[:20]


@needs_soffice
def test_pdf_to_word_flow(fx, ctx):
    import docx

    r = convert.to_word(ctx, fx["financial"], out("fin_flow.docx"), mode="flow")
    d = docx.Document(r["output"])
    assert r["tables"] >= 1
    allcells = [c.text for t in d.tables for row in t.rows for c in row.cells]
    assert "(45,000.00)" in allcells


# ---- PDF -> Excel ----------------------------------------------------------------

@needs_soffice
def test_pdf_to_excel_financial_values(fx, ctx):
    import openpyxl

    r = convert.to_excel(ctx, fx["financial"], out("fin.xlsx"))
    wb = openpyxl.load_workbook(r["output"])
    assert "Extraction Report" in wb.sheetnames
    values = []
    for ws in wb.worksheets[1:]:
        for row in ws.iter_rows(values_only=True):
            values.append(row)
    flat = [v for row in values for v in row if v is not None]
    assert 123456.78 in flat
    assert -45000.0 in flat           # parentheses -> negative, sign preserved
    assert "000987654321" in flat      # reference with leading zeros stays text (and is not split by the cell wrap)
    assert "412345678901" in flat      # long reference stays text
    assert "NEFT-HDFC0001234-ACME SUPPLIES" in flat  # wrapped cell rebuilt without inserting spaces
    assert "UPI/CR/412345678901/Client A" in flat
    other = wb["Other text"]
    assert any("000123456789" in str(c.value) for row in other.iter_rows() for c in row)
    assert "27ABCDE1234F1Z5" in " ".join(str(v) for v in flat)
    # Bank statement table spans pages and must be merged into one sheet with every row.
    bank = [ws for ws in wb.worksheets if any(c.value == "Description" for c in ws[1])]
    assert bank, wb.sheetnames
    rows = list(bank[0].iter_rows(values_only=True))
    assert len(rows) == 1 + 18 * len(BANK_ROWS), len(rows)
    import datetime as dt

    assert any(isinstance(v, dt.datetime) and v.date() == dt.date(2025, 4, 15) for row in rows for v in row)
    rep = r["report"]
    assert all(p["coverage"] is None or p["coverage"] == 100 for p in rep), rep


@needs_soffice
def test_pdf_to_csv(fx, ctx):
    r = convert.to_excel(ctx, fx["financial"], out("fin.csv"))
    text = open(r["output"], encoding="utf-8-sig").read()
    assert "1,23,456.78" in text and ",000987654321," in text and "NEFT-HDFC0001234-ACME SUPPLIES" in text


# ---- PDF -> PowerPoint ------------------------------------------------------------

@needs_soffice
@pytest.mark.parametrize("mode", ["image", "hybrid", "editable"])
def test_pdf_to_ppt_modes(fx, ctx, mode):
    import pptx

    r = convert.to_powerpoint(ctx, fx["financial"], out(f"fin_{mode}.pptx"), mode=mode, pages="1")
    p = pptx.Presentation(r["output"])
    shapes = list(p.slides[0].shapes)
    texts = " ".join(s.text_frame.text for s in shapes if s.has_text_frame)
    if mode == "image":
        assert r["textBoxes"] == 0 and r["pictures"] == 1 and texts.strip() == ""
    else:
        assert "ACME Traders" in texts and "1,23,456.78" in texts


# ---- PDF -> images / text / html / md ------------------------------------------------

def test_pdf_to_images(fx, ctx):
    from PIL import Image

    od = os.path.join(os.path.dirname(out("x")), "imgs")
    r = convert.to_images(ctx, fx["c"], od, format="png", dpi=100, pages="2-3", overwrite=True)
    assert r["count"] == 2
    for o, pg in zip(r["outputs"], (2, 3)):
        assert o["output"].endswith(f"c_page_{pg:03d}.png")
        with Image.open(o["output"]) as im:
            assert im.format == "PNG"
            assert abs(im.size[0] - 595 * 100 / 72) <= 1 and abs(im.size[1] - 842 * 100 / 72) <= 1, im.size
            assert im.size == (o["width"], o["height"])
    for fmt, pil in (("jpg", "JPEG"), ("tiff", "TIFF"), ("bmp", "BMP"), ("webp", "WEBP")):
        r = convert.to_images(ctx, fx["a"], od, format=fmt, dpi=72, pages="1", overwrite=True)
        with Image.open(r["outputs"][0]["output"]) as im:
            assert im.format == pil


def test_pdf_to_text_layout_html_md(fx, ctx):
    r = convert.to_text(ctx, fx["twenty"], out("t.txt"), pages="1-2")
    t = open(r["output"], encoding="utf-8").read()
    assert "Report page 1" in t and "Report page 2" in t
    r = convert.to_html(ctx, fx["a"], out("t.html"))
    assert "A page 1" in open(r["output"], encoding="utf-8").read()
    r = convert.to_markdown(ctx, fx["a"], out("t.md"))
    assert "A page 1" in open(r["output"], encoding="utf-8").read()


# ---- into PDF -------------------------------------------------------------------------

@needs_soffice
def test_office_to_pdf_word_excel_ppt(fx, ctx):
    import openpyxl
    import pptx

    r = convert.office_to_pdf(ctx, fx["financial_docx"], out("docx2pdf.pdf"))
    assert "ACME Traders" in page_texts(r["output"])[0]
    xl = out("sheet.xlsx")
    wb = openpyxl.Workbook()
    wb.active.append(["Item", "Amount"])
    wb.active.append(["Audit fee", 75000])
    wb.save(xl)
    r = convert.office_to_pdf(ctx, xl, out("xlsx2pdf.pdf"), singlePageSheets=True)
    assert "Audit fee" in page_texts(r["output"])[0]
    pp = out("deck.pptx")
    prs = pptx.Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Board report Q1"
    prs.save(pp)
    r = convert.office_to_pdf(ctx, pp, out("pptx2pdf.pdf"))
    assert "Board report Q1" in page_texts(r["output"])[0]


def test_images_to_pdf(fx, ctx):
    r = create.from_images(ctx, [fx["jpg"], fx["png"]], out("imgs.pdf"), pageSize="A4", margin=36)
    d = pymupdf.open(r["output"])
    assert d.page_count == 2
    assert round(d[0].rect.width) == 842  # landscape image -> landscape page (auto)
    d.close()
    r = create.from_images(ctx, [fx["jpg"]], out("imgs_fit.pdf"), pageSize="fit")
    assert r["pageCount"] == 1


def test_text_to_pdf(fx, ctx):
    r = create.from_text(ctx, fx["txt"], out("txt.pdf"))
    t = "".join(page_texts(r["output"]))
    assert "Line 119: amount ₹ 119,000" in t


@needs_soffice
def test_assemble_mixed(fx, ctx):
    r = create.assemble(ctx, [fx["a"], fx["jpg"], fx["financial_docx"], fx["txt"]], out("assembled.pdf"))
    d = pymupdf.open(r["output"])
    assert d.page_count >= 5
    assert [t[1] for t in d.get_toc() if t[0] == 1][:2] == ["a", "photo"]
    d.close()


# ---- OCR --------------------------------------------------------------------------------

@needs_tesseract
def test_ocr_creates_searchable_pdf(fx, ctx):
    src = fx["scanned"]
    assert not any(t.strip() for t in page_texts(src))
    r = ocr.run(ctx, src, out("scanned_ocr.pdf"), languages=["eng"], pages="1")
    assert r["processed"] == [1]
    d = pymupdf.open(r["output"])
    assert d[0].search_for("Traders"), "OCR text not searchable"
    assert d[0].search_for("Revenue")
    # The original scan image is preserved untouched.
    assert len(d[0].get_images()) == 1
    d.close()
    assert r["pages"][1]["meanConfidence"] is not None


@needs_tesseract
def test_ocr_skips_text_pages(fx, ctx):
    from nayanpdf_engine.errors import EngineError

    with pytest.raises(EngineError):
        ocr.run(ctx, fx["a"], out("a_ocr.pdf"))


@needs_tesseract
@needs_soffice
def test_scanned_to_excel_with_ocr(fx, ctx):
    import openpyxl

    r = convert.to_excel(ctx, fx["scanned"], out("scan.xlsx"), pages="1", ocr="auto")
    assert r.get("ocr")
    wb = openpyxl.load_workbook(r["output"])
    flat = [str(v) for ws in wb.worksheets[1:] for row in ws.iter_rows(values_only=True) for v in row if v is not None]
    assert any("Revenue" in v for v in flat)


@needs_tesseract
def test_ocr_plain_text_extraction(fx, ctx):
    r = ocr.text(ctx, fx["scanned"], out("scanned_ocr.txt"), pages="1")
    txt = open(r["output"], encoding="utf-8").read()
    assert "Revenue" in txt


@needs_tesseract
def test_ocr_turns_sideways_scan_upright(fx, ctx):
    from nayanpdf_engine.util import tools
    if not os.path.exists(os.path.join(tools.tessdata_dir() or "", "osd.traineddata")):
        pytest.skip("orientation data (osd) not installed")
    # A scan fed into the scanner sideways: the page image is rotated 90 degrees.
    from PIL import Image
    src = pymupdf.open(fx["scanned"])
    pix = src[0].get_pixmap(dpi=200)
    src.close()
    img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples).rotate(90, expand=True)
    png = out("sideways.png")
    img.save(png)
    side = pymupdf.open()
    pg = side.new_page(width=img.width * 72 / 200, height=img.height * 72 / 200)
    pg.insert_image(pg.rect, filename=png)
    side_pdf = out("sideways.pdf")
    side.save(side_pdf)
    side.close()
    r = ocr.run(ctx, side_pdf, out("sideways_ocr.pdf"), languages=["eng"])
    assert r.get("rotated"), r
    d = pymupdf.open(r["output"])
    assert d[0].search_for("Revenue"), "text on the corrected page is not searchable"
    d.close()


@needs_soffice
def test_office_conversion_never_fetches_linked_content(ctx, tmp_path, monkeypatch):
    """Office files can reference images or data by URL. Converting them must not
    contact the network: a local HTTP server stands in for the remote host and
    must receive no request."""
    import http.server
    import io
    import threading

    import docx
    import openpyxl
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from PIL import Image

    hits = []

    class Handler(http.server.BaseHTTPRequestHandler):
        def do_GET(self):  # noqa: N802
            hits.append(self.path)
            buf = io.BytesIO()
            Image.new("RGB", (40, 40), (200, 0, 0)).save(buf, "PNG")
            self.send_response(200)
            self.send_header("Content-Type", "image/png")
            self.end_headers()
            self.wfile.write(buf.getvalue())

        def log_message(self, *args):
            pass

    srv = http.server.HTTPServer(("127.0.0.1", 0), Handler)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    base = f"http://127.0.0.1:{srv.server_port}"
    monkeypatch.setenv("NAYANPDF_LO_PROFILES", str(tmp_path / "profiles"))
    try:
        # Word file whose picture is linked to a URL instead of embedded.
        d = docx.Document()
        d.add_paragraph("Linked image test")
        buf = io.BytesIO()
        Image.new("RGB", (40, 40), (0, 0, 200)).save(buf, "PNG")
        buf.seek(0)
        d.add_picture(buf)
        rid = d.part.relate_to(base + "/tracker.png", RT.IMAGE, is_external=True)
        ns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
        for blip in d.element.body.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}blip"):
            blip.attrib.pop(ns + "embed", None)
            blip.set(ns + "link", rid)
        src_docx = str(tmp_path / "linked.docx")
        d.save(src_docx)
        r = convert.office_to_pdf(ctx, src_docx, str(tmp_path / "linked.pdf"))
        assert "Linked image test" in page_texts(r["output"])[0]

        # Excel file with a formula that would download data when recalculated.
        wb = openpyxl.Workbook()
        wb.active["A1"] = "Rate"
        wb.active["B1"] = f'=WEBSERVICE("{base}/rate")'
        src_xlsx = str(tmp_path / "webservice.xlsx")
        wb.save(src_xlsx)
        r = convert.office_to_pdf(ctx, src_xlsx, str(tmp_path / "webservice.pdf"))
        assert "Rate" in page_texts(r["output"])[0]
    finally:
        srv.shutdown()
    assert hits == [], f"conversion contacted the network: {hits}"
