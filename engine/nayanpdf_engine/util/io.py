"""Safe file output.

Every output is produced in the job's private temp directory, validated, and
only then moved to its destination. If the destination is on the same volume
the move is an atomic rename; otherwise the file is copied to a hidden
``.part`` file beside the destination and renamed. Existing files are never
overwritten unless the caller explicitly allows it, and source files are
never touched.
"""

from __future__ import annotations

import errno
import os
import re
import shutil
import sys
import zipfile
from typing import Callable

from ..errors import (
    ACCESS_DENIED,
    FILE_NOT_FOUND,
    NO_SPACE,
    OPERATION_FAILED,
    OUTPUT_EXISTS,
    VALIDATION_FAILED,
    EngineError,
    invalid,
)

IS_WIN = sys.platform == "win32"
_INVALID_NAME_CHARS = re.compile(r'[<>:"/\\|?*\x00-\x1f]')
_RESERVED = {"CON", "PRN", "AUX", "NUL", *(f"COM{i}" for i in range(1, 10)), *(f"LPT{i}" for i in range(1, 10))}


def os_path(p: str) -> str:
    """Return a path usable by the OS even when it exceeds MAX_PATH on Windows."""
    if not IS_WIN:
        return p
    p = os.path.abspath(p)
    if len(p) >= 240 and not p.startswith("\\\\?\\"):
        if p.startswith("\\\\"):
            return "\\\\?\\UNC\\" + p[2:]
        return "\\\\?\\" + p
    return p


def os_error_to_engine(e: OSError) -> EngineError:
    if e.errno == errno.ENOSPC or getattr(e, "winerror", None) in (39, 112):
        return EngineError(NO_SPACE, "There is not enough free disk space to write the output.")
    if e.errno in (errno.EACCES, errno.EPERM) or getattr(e, "winerror", None) in (5, 32, 33):
        return EngineError(ACCESS_DENIED, "Access denied. The file may be open in another program or the folder is read-only.",
                           {"path": getattr(e, "filename", None)})
    if e.errno == errno.ENOENT:
        return EngineError(FILE_NOT_FOUND, "File or folder not found.", {"path": getattr(e, "filename", None)})
    return EngineError(OPERATION_FAILED, f"File system error: {e.strerror or e}")


def require_file(path: str) -> str:
    if not path or not isinstance(path, str):
        raise invalid("A file path is required.")
    if not os.path.isfile(os_path(path)):
        raise EngineError(FILE_NOT_FOUND, f"File not found: {os.path.basename(path)}", {"path": path})
    return path


def safe_filename(name: str, fallback: str = "output") -> str:
    name = _INVALID_NAME_CHARS.sub("_", name).strip().rstrip(".")
    if not name:
        name = fallback
    stem = name.split(".")[0].upper()
    if stem in _RESERVED:
        name = "_" + name
    return name[:200]


def unique_path(path: str) -> str:
    """'a.pdf' -> 'a (1).pdf' -> 'a (2).pdf' ... for a path that does not exist."""
    if not os.path.exists(os_path(path)):
        return path
    base, ext = os.path.splitext(path)
    m = re.match(r"^(.*) \((\d+)\)$", base)
    if m:
        base = m.group(1)
    i = 1
    while True:
        cand = f"{base} ({i}){ext}"
        if not os.path.exists(os_path(cand)):
            return cand
        i += 1


def derive_output(source: str, output_dir: str | None, suffix: str, ext: str,
                  name: str | None = None) -> str:
    """Build an output path like <dir>/<source stem><suffix><ext>."""
    stem = os.path.splitext(os.path.basename(source))[0]
    out_dir = output_dir or os.path.dirname(os.path.abspath(source))
    filename = safe_filename(name if name else f"{stem}{suffix}")
    if not filename.lower().endswith(ext.lower()):
        filename += ext
    return os.path.join(out_dir, filename)


def apply_name_template(template: str, source: str, **tokens) -> str:
    """Expand a naming rule such as '{name}_compressed' or '{name}_page_{page:03}'."""
    stem = os.path.splitext(os.path.basename(source))[0]
    values = {"name": stem, **tokens}
    try:
        out = template.format(**values)
    except (KeyError, IndexError, ValueError) as e:
        raise invalid(f"Invalid naming rule '{template}': {e}")
    return safe_filename(out)


def same_file(a: str, b: str) -> bool:
    try:
        return os.path.samefile(os_path(a), os_path(b))
    except OSError:
        return os.path.normcase(os.path.abspath(a)) == os.path.normcase(os.path.abspath(b))


class OutputFile:
    """A pending output file.

    Usage:
        out = OutputFile(ctx, "/x/report.pdf", overwrite=False, sources=[src])
        doc.save(out.temp)
        final = out.commit(validate_pdf)
    """

    def __init__(self, ctx, final_path: str, overwrite: bool = False, sources: list[str] | None = None):
        if not final_path:
            raise invalid("An output path is required.")
        final_path = os.path.abspath(final_path)
        self.ctx = ctx
        self.overwrite = bool(overwrite)
        self.sources = [s for s in (sources or []) if s]
        for s in self.sources:
            if os.path.exists(os_path(final_path)) and same_file(s, final_path) and not self.overwrite:
                raise EngineError(OUTPUT_EXISTS, "The output would replace the source document. Choose a different name or explicitly allow overwriting.")
        out_dir = os.path.dirname(final_path)
        if not os.path.isdir(os_path(out_dir)):
            try:
                os.makedirs(os_path(out_dir), exist_ok=True)
            except OSError as e:
                raise os_error_to_engine(e)
        self.requested = final_path
        ext = os.path.splitext(final_path)[1]
        self.temp = ctx.tmp_path(suffix=ext or ".tmp")
        self.final: str | None = None

    def commit(self, validate: Callable[[str], dict | None] | None = None) -> str:
        if not os.path.isfile(self.temp) or os.path.getsize(self.temp) == 0:
            raise EngineError(VALIDATION_FAILED, "The operation did not produce an output file.")
        info = None
        if validate:
            try:
                info = validate(self.temp)
            except EngineError:
                raise
            except Exception as e:  # noqa: BLE001
                raise EngineError(VALIDATION_FAILED, f"The generated file failed validation: {e}")
        self.validation = info or {}
        target = self.requested
        if os.path.exists(os_path(target)):
            if not self.overwrite:
                target = unique_path(target)
        self._move_into_place(target)
        self.final = target
        return target

    def _move_into_place(self, target: str) -> None:
        size = os.path.getsize(self.temp)
        target_dir = os.path.dirname(target)
        try:
            free = shutil.disk_usage(os_path(target_dir)).free
            if free < size + 1024 * 1024:
                raise EngineError(NO_SPACE, "Not enough free disk space at the destination.",
                                  {"required": size, "free": free})
        except OSError:
            pass
        try:
            os.replace(self.temp, os_path(target))
            return
        except OSError as e:
            if e.errno not in (errno.EXDEV,) and getattr(e, "winerror", None) != 17:
                if e.errno in (errno.EACCES, errno.EPERM, errno.ENOSPC):
                    raise os_error_to_engine(e)
        # Different volume: copy beside the target then rename.
        part = os.path.join(target_dir, "." + os.path.basename(target) + ".nayan-part")
        self.ctx.partial(part)
        try:
            shutil.copyfile(self.temp, os_path(part))
            os.replace(os_path(part), os_path(target))
        except OSError as e:
            try:
                if os.path.exists(os_path(part)):
                    os.remove(os_path(part))
            except OSError:
                pass
            raise os_error_to_engine(e)
        finally:
            try:
                os.remove(self.temp)
            except OSError:
                pass


# ---------------------------------------------------------------------------
# Validators: reopen generated files before reporting success.
# ---------------------------------------------------------------------------

def validate_pdf(expected_pages: int | None = None, password: str | None = None):
    def _v(path: str) -> dict:
        import pymupdf

        try:
            doc = pymupdf.open(path)
        except Exception as e:  # noqa: BLE001
            raise EngineError(VALIDATION_FAILED, f"Generated PDF cannot be reopened: {e}")
        try:
            if doc.needs_pass:
                if not password or not doc.authenticate(password):
                    # Encrypted output: verify it at least opens structurally.
                    return {"pages": doc.page_count, "encrypted": True}
            n = doc.page_count
            if n <= 0:
                raise EngineError(VALIDATION_FAILED, "Generated PDF has no pages.")
            if expected_pages is not None and n != expected_pages:
                raise EngineError(VALIDATION_FAILED,
                                  f"Generated PDF has {n} pages; expected {expected_pages}.")
            # Touch first and last page to make sure the page tree resolves.
            doc[0].bound()
            doc[n - 1].bound()
            return {"pages": n}
        finally:
            doc.close()

    return _v


def validate_docx(path: str) -> dict:
    import docx

    if not zipfile.is_zipfile(path):
        raise EngineError(VALIDATION_FAILED, "Generated DOCX is not a valid package.")
    d = docx.Document(path)
    paras = sum(1 for p in d.paragraphs if p.text.strip())
    return {"paragraphs": paras, "tables": len(d.tables)}


def validate_xlsx(path: str) -> dict:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        sheets = wb.sheetnames
        if not sheets:
            raise EngineError(VALIDATION_FAILED, "Generated workbook has no sheets.")
        return {"sheets": len(sheets)}
    finally:
        wb.close()


def validate_pptx(path: str) -> dict:
    import pptx

    p = pptx.Presentation(path)
    n = len(p.slides)
    if n == 0:
        raise EngineError(VALIDATION_FAILED, "Generated presentation has no slides.")
    return {"slides": n}


def validate_image(path: str) -> dict:
    from PIL import Image

    with Image.open(path) as im:
        im.verify()
    with Image.open(path) as im:
        return {"width": im.width, "height": im.height, "format": im.format}


def validate_nonempty(path: str) -> dict:
    size = os.path.getsize(path)
    return {"bytes": size}


def file_size(path: str) -> int:
    return os.path.getsize(os_path(path))
