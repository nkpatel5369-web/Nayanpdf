"""Conversions from PDF to other formats, and Office documents to PDF.

Every converter writes a genuine file in the target format and reopens it
before success is reported. Nothing is produced by renaming files.
"""

from __future__ import annotations

import csv
import io
import logging
import os
import re
import shutil
import statistics

import pymupdf

from ..errors import EngineError, TOOL_MISSING, UNSUPPORTED_FORMAT, VALIDATION_FAILED, invalid
from ..registry import method
from ..util import tools
from ..util.io import (
    OutputFile,
    apply_name_template,
    derive_output,
    os_path,
    require_file,
    validate_docx,
    validate_image,
    validate_nonempty,
    validate_pdf,
    validate_pptx,
    validate_xlsx,
)
from ..util.numbers import digits_of, numeric_tokens, parse_date, parse_number
from ..util.pdf import int_to_rgb, open_pdf, parse_pages, rgb_to_hex


def _maybe_ocr(ctx, path, password, ocr, languages, require_any_text=True):
    """Return (path_to_use, password_to_use, ocr_info) - OCR'd temp copy when needed."""
    if ocr == "off":
        return path, password, None
    doc = open_pdf(path, password, require=["copy"])
    textless = [i for i in range(doc.page_count) if not doc[i].get_text("text").strip()]
    doc.close()
    if ocr == "auto" and not textless:
        return path, password, None
    from .ocr import _check_langs, searchable_copy

    lang = _check_langs(languages or ["eng"])
    ctx.status("Running OCR on scanned pages")
    tmp, info = searchable_copy(ctx, path, password, lang)
    review = [p for p, st in info["pages"].items() if st["review"]]
    if review:
        ctx.warn(f"OCR confidence is low on page(s) {', '.join(map(str, review[:20]))}; review the converted content.")
    ctx.warn("Scanned pages were recognised with OCR. Check figures and identifiers against the original.")
    return tmp, None, {"ocrPages": info["processed"], "reviewPages": review}


# ============================================================================
# PDF -> Word
# ============================================================================

class _Pdf2DocxProgress(logging.Handler):
    def __init__(self, ctx, n):
        super().__init__(logging.INFO)
        self.ctx = ctx
        self.n = n
        self.phase = 0
        self.errors: list[str] = []

    def emit(self, record):
        msg = record.getMessage()
        if "[3/4]" in msg:
            self.phase = 0
        elif "[4/4]" in msg:
            self.phase = 1
        m = re.match(r"\((\d+)/(\d+)\) Page (\d+)", msg)
        if m:
            i, n = int(m.group(1)), int(m.group(2))
            self.ctx.progress(self.phase * n + i, 2 * n, ("Analysing" if self.phase == 0 else "Writing") + f" page {m.group(3)}")
        if record.levelno >= logging.ERROR:
            pm = re.search(r"page (\d+)", msg)
            self.errors.append(pm.group(1) if pm else msg[:120])


@method("convert.toWord")
def to_word(ctx, path: str, output: str, mode: str = "layout", pages=None, ocr: str = "auto",
            languages: list | None = None, password: str | None = None, overwrite: bool = False,
            includeImages: bool = True):
    """mode='layout' keeps positioning (pdf2docx); mode='flow' favours easy editing."""
    src, pw, ocr_info = _maybe_ocr(ctx, path, password, ocr, languages)
    probe = open_pdf(src, pw, require=["copy"])
    sel = parse_pages(pages, probe.page_count)
    probe.close()
    if ocr_info and mode == "layout":
        ctx.warn("Scanned pages are converted in 'editable text' mode because layout reconstruction of OCR text is unreliable.")
        mode = "flow"
    out = OutputFile(ctx, output, overwrite, sources=[path])
    if mode == "layout":
        from pdf2docx import Converter

        handler = _Pdf2DocxProgress(ctx, len(sel))
        root = logging.getLogger()
        root.addHandler(handler)
        try:
            cv = Converter(src, password=pw)
            try:
                cv.convert(out.temp, pages=sel, multi_processing=False)
            finally:
                cv.close()
        except Exception as e:  # noqa: BLE001
            raise EngineError("OPERATION_FAILED", f"Layout conversion failed: {e}. Try 'Editable text' mode.")
        finally:
            root.removeHandler(handler)
        if handler.errors:
            ctx.warn("Some pages could not be reconstructed and were skipped: " + ", ".join(handler.errors[:10]) +
                     ". Try 'Editable text' mode for those pages.")
    elif mode == "flow":
        _flow_docx(ctx, src, pw, sel, out.temp, includeImages)
    else:
        raise invalid("mode must be 'layout' or 'flow'.")
    final = out.commit(validate_docx)
    info = out.validation
    if info.get("paragraphs", 0) == 0 and info.get("tables", 0) == 0:
        ctx.warn("The Word document contains no text. The PDF may be scanned without OCR, or text is drawn as graphics.")
    res = {"output": final, "mode": mode, **info}
    if ocr_info:
        res["ocr"] = ocr_info
    return res


def _flow_docx(ctx, path, password, sel, target, include_images):
    import docx
    from docx.enum.section import WD_ORIENT
    from docx.shared import Pt, RGBColor, Inches

    doc = open_pdf(path, password, require=["copy"])
    d = docx.Document()
    style = d.styles["Normal"]
    style.font.size = Pt(10.5)
    # Body text size = most common span size.
    sizes = []
    for i in sel[:30]:
        for b in doc[i].get_text("dict")["blocks"]:
            for ln in b.get("lines", []):
                for s in ln["spans"]:
                    if s["text"].strip():
                        sizes.append(round(s["size"]))
    body = statistics.mode(sizes) if sizes else 11
    first = doc[sel[0]].rect if sel else None
    if first is not None and first.width > first.height:
        sec = d.sections[0]
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for k, i in enumerate(sel):
        ctx.progress(k, len(sel), f"Page {i + 1}")
        page = doc[i]
        tables = []
        try:
            tf = page.find_tables()
            tables = [t for t in tf.tables if t.row_count >= 2 and t.col_count >= 2]
        except Exception:  # noqa: BLE001
            tables = []
        tboxes = [pymupdf.Rect(t.bbox) for t in tables]
        items = []
        for t in tables:
            items.append((t.bbox[1], "table", t))
        data = page.get_text("dict", sort=True)
        for b in data["blocks"]:
            bb = pymupdf.Rect(b["bbox"])
            if any((bb & tb).get_area() > 0.5 * bb.get_area() for tb in tboxes):
                continue
            if b["type"] == 1:
                if include_images:
                    items.append((bb.y0, "image", b))
                continue
            items.append((bb.y0, "text", b))
        items.sort(key=lambda x: x[0])
        for _, kind, obj in items:
            if kind == "table":
                rows = obj.extract()
                ncols = max(len(r) for r in rows) if rows else 0
                if not ncols:
                    continue
                tbl = d.add_table(rows=len(rows), cols=ncols)
                tbl.style = "Table Grid"
                for ri, row in enumerate(rows):
                    for ci in range(ncols):
                        v = row[ci] if ci < len(row) else None
                        tbl.cell(ri, ci).text = (v or "").replace("\n", " ")
                d.add_paragraph()
            elif kind == "image":
                try:
                    pix = pymupdf.Pixmap(obj["image"]) if isinstance(obj.get("image"), (bytes, bytearray)) else None
                    if pix is not None:
                        if pix.alpha or (pix.colorspace and pix.colorspace.n > 3):
                            pix = pymupdf.Pixmap(pymupdf.csRGB, pix)
                        w_in = min(6.5, (obj["bbox"][2] - obj["bbox"][0]) / 72)
                        d.add_picture(io.BytesIO(pix.tobytes("png")), width=Inches(max(0.3, w_in)))
                except Exception:  # noqa: BLE001
                    pass
            else:
                lines = obj["lines"]
                spans = [s for ln in lines for s in ln["spans"] if s["text"].strip()]
                if not spans:
                    continue
                size = max(s["size"] for s in spans)
                bold = all((s["flags"] & 16) or "bold" in s["font"].lower() for s in spans)
                text_lines = ["".join(s["text"] for s in ln["spans"]).strip() for ln in lines]
                text = " ".join(t for t in text_lines if t)
                if size >= body * 1.45 and len(text) < 200:
                    d.add_heading(text, level=1)
                    continue
                if size >= body * 1.2 and len(text) < 200:
                    d.add_heading(text, level=2)
                    continue
                if bold and len(text) < 120 and size >= body:
                    d.add_heading(text, level=3)
                    continue
                p = d.add_paragraph()
                for li, ln in enumerate(lines):
                    for s in ln["spans"]:
                        r = p.add_run(s["text"])
                        r.font.size = Pt(max(6, round(s["size"] * 2) / 2))
                        r.bold = bool(s["flags"] & 16) or "bold" in s["font"].lower()
                        r.italic = bool(s["flags"] & 2) or "italic" in s["font"].lower()
                        col = int_to_rgb(s["color"])
                        if any(c > 0.05 for c in col):
                            r.font.color.rgb = RGBColor(*(int(c * 255) for c in col))
                    if li < len(lines) - 1:
                        p.add_run(" ")
        if k < len(sel) - 1:
            d.add_page_break()
    doc.close()
    d.save(target)


# ============================================================================
# PDF -> Excel / CSV
# ============================================================================

def _clean_cell(v):
    if v is None:
        return None
    return re.sub(r"[ \t]*\n[ \t]*", " ", str(v)).strip()


class _PageChars:
    """Raw characters of a page grouped by text line, for exact cell rebuilding."""

    def __init__(self, page: pymupdf.Page):
        self.lines = []  # (line_bbox, [(char, bbox)])
        raw = page.get_text("rawdict", flags=pymupdf.TEXTFLAGS_RAWDICT & ~pymupdf.TEXT_PRESERVE_IMAGES)
        wrap_trailing = 0
        for b in raw["blocks"]:
            blines = b.get("lines", [])
            for k, ln in enumerate(blines):
                chars = [(c["c"], pymupdf.Rect(c["bbox"])) for s in ln["spans"] for c in s["chars"]]
                if not chars:
                    continue
                self.lines.append((pymupdf.Rect(ln["bbox"]), chars))
                if k + 1 < len(blines) and chars[-1][0] == " ":
                    wrap_trailing += 1
        # True if this PDF keeps the space at soft line wraps (then a wrap with no
        # trailing space is a break inside a word/number and must be joined without a space).
        self.keeps_wrap_spaces = wrap_trailing > 0

    def cell_text(self, bbox) -> tuple[str, bool, bool]:
        """Return (text, ambiguous_join, ended_with_space)."""
        r = pymupdf.Rect(bbox)
        parts: list[str] = []
        for lb, chars in self.lines:
            if not lb.intersects(r):
                continue
            seg = "".join(c for c, cb in chars if r.contains(pymupdf.Point((cb.x0 + cb.x1) / 2, (cb.y0 + cb.y1) / 2)))
            if seg.strip():
                parts.append(seg)
        if not parts:
            return "", False, False
        out = parts[0]
        ambiguous = False
        for nxt in parts[1:]:
            prev = out
            if prev.endswith((" ", "\t")):
                out = prev.rstrip() + " " + nxt.lstrip()
            elif prev.endswith(("-", "/", "\\", "_")):
                out = prev + nxt.lstrip()
            elif self.keeps_wrap_spaces:
                out = prev + nxt.lstrip()
            else:
                if prev[-1:].isalnum() and nxt.lstrip()[:1].isalnum():
                    ambiguous = True
                out = prev + " " + nxt.lstrip()
        return re.sub(r"[ \t]+", " ", out).strip(), ambiguous, parts[-1].endswith((" ", "\t"))


def join_wrapped(prev: str, nxt: str, prev_trailing_space: bool, keeps_wrap_spaces: bool) -> tuple[str, bool]:
    """Join two fragments of one cell that were split by a line or page break."""
    if not prev:
        return nxt, False
    if not nxt:
        return prev, False
    if prev_trailing_space:
        return prev + " " + nxt, False
    if prev.endswith(("-", "/", "\\", "_")):
        return prev + nxt, False
    if keeps_wrap_spaces:
        return prev + nxt, False
    return prev + " " + nxt, prev[-1:].isalnum() and nxt[:1].isalnum()


def _table_rows(page_chars: _PageChars, table) -> tuple[list[list], int, list[list]]:
    """Rebuild table rows from raw characters; returns (rows, ambiguous_cells, trailing_space_flags)."""
    rows = []
    trail = []
    amb = 0
    for row in table.rows:
        out_row = []
        tr_row = []
        for cb in row.cells:
            if cb is None:
                out_row.append(None)
                tr_row.append(False)
                continue
            txt, a, tr = page_chars.cell_text(cb)
            amb += int(a)
            out_row.append(txt)
            tr_row.append(tr)
        rows.append(out_row)
        trail.append(tr_row)
    if not any(any(c for c in r) for r in rows):
        fallback = [[_clean_cell(c) for c in r] for r in table.extract()]
        return fallback, 0, [[False] * len(r) for r in fallback]
    return rows, amb, trail


def _extract_tables(page, strategy: str):
    tried = []
    order = {"auto": ["lines", "text"], "lines": ["lines"], "text": ["text"]}.get(strategy, ["lines", "text"])
    for strat in order:
        try:
            tf = page.find_tables(strategy=strat) if strat != "lines" else page.find_tables()
        except Exception:  # noqa: BLE001
            continue
        tabs = [t for t in tf.tables if t.row_count >= 2 and t.col_count >= 2]
        tried.append(strat)
        if tabs:
            return tabs, strat
    return [], "+".join(tried) or strategy


@method("convert.toExcel")
def to_excel(ctx, path: str, output: str, pages=None, strategy: str = "auto", layout: str = "sheetPerTable",
             mergeAcrossPages: bool = True, convertNumbers: bool = True, dateOrder: str = "DMY",
             ocr: str = "auto", languages: list | None = None, password: str | None = None,
             overwrite: bool = False, csvMode: str = "combined"):
    """Extract tables into XLSX (or CSV when output ends with .csv)."""
    is_csv = output.lower().endswith(".csv")
    src, pw, ocr_info = _maybe_ocr(ctx, path, password, ocr, languages)
    doc = open_pdf(src, pw, require=["copy"])
    sel = parse_pages(pages, doc.page_count)
    tables = []  # dicts: page, rows, strategy, bbox
    report = []
    other_text: list = []
    for k, i in enumerate(sel):
        ctx.progress(k, len(sel), f"Detecting tables on page {i + 1}")
        page = doc[i]
        page_text = page.get_text("text")
        tabs, strat = _extract_tables(page, strategy)
        page_tables = []
        pchars = _PageChars(page) if tabs else None
        ambiguous_cells = 0
        for t in tabs:
            rows, amb, trail = _table_rows(pchars, t)
            ambiguous_cells += amb
            keep = [k for k, r in enumerate(rows) if any(c for c in r)]
            rows = [rows[k] for k in keep]
            trail = [trail[k] for k in keep]
            if not rows:
                continue
            header_names = None
            try:
                if t.header and t.header.external and t.header.names:
                    header_names = [_clean_cell(c) for c in t.header.names]
            except Exception:  # noqa: BLE001
                header_names = None
            if header_names and any(header_names) and rows[0] != header_names:
                rows.insert(0, header_names)
                trail.insert(0, [False] * len(header_names))
            page_tables.append({"page": i + 1, "rows": rows, "trail": trail, "strategy": strat,
                                "ncols": max(len(r) for r in rows), "keepsWrap": pchars.keeps_wrap_spaces})
        tokens_page = numeric_tokens(page_text)
        captured_texts: list[str] = []
        for t in page_tables:
            for r in t["rows"]:
                captured_texts.extend(c for c in r if c)
        if not page_tables and page_text.strip():
            # Keep the content: one row per text line, flagged for review.
            lines = [[ln] for ln in page_text.splitlines() if ln.strip()]
            page_tables.append({"page": i + 1, "rows": lines, "strategy": "lines-as-rows", "ncols": 1, "noTable": True})
            captured_texts = [page_text]
        # Text outside tables is kept on an "Other text" sheet so nothing is silently dropped.
        tboxes = [pymupdf.Rect(t.bbox) for t in tabs]
        other = []
        if tabs:
            for b in page.get_text("dict", sort=True)["blocks"]:
                for ln in b.get("lines", []):
                    lb = pymupdf.Rect(ln["bbox"])
                    if any((lb & tb).get_area() > 0.5 * max(lb.get_area(), 1e-6) for tb in tboxes):
                        continue
                    lt = "".join(s_["text"] for s_ in ln["spans"]).strip()
                    if lt:
                        other.append(lt)
                        captured_texts.append(lt)
        if other:
            other_text.append((i + 1, other))
        # Coverage: every digit printed on the page must appear in the workbook
        # (robust to line wrapping); plus a token-level sample of what is missing.
        from collections import Counter

        page_digits = Counter(ch for ch in page_text if ch.isdigit())
        cap_digits = Counter(ch for txt in captured_texts for ch in txt if ch.isdigit())
        total_digits = sum(page_digits.values())
        missing_digits = sum((page_digits - cap_digits).values())
        cov = round(100.0 * (1 - missing_digits / total_digits), 1) if total_digits else None
        cap_strings = [digits_of(t) for t in captured_texts]
        missing = [tok for tok in tokens_page if not any(tok in c for c in cap_strings)]
        report.append({"page": i + 1, "tables": len([t for t in page_tables if not t.get("noTable")]),
                       "ambiguousCells": ambiguous_cells,
                       "strategy": page_tables[0]["strategy"] if page_tables else strat,
                       "numbersOnPage": len(tokens_page), "numbersCaptured": len(tokens_page) - len(missing),
                       "coverage": cov, "missingSample": missing[:8],
                       "review": (cov is not None and cov < 100.0) or ambiguous_cells > 0
                                 or any(t.get("noTable") for t in page_tables) or not page_tables})
        tables.extend(page_tables)
    doc.close()
    if not tables:
        raise EngineError("NOT_SUPPORTED", "No text or tables were found. If the PDF is scanned, enable OCR.")

    # Merge tables continuing across pages (same column count, consecutive pages).
    if mergeAcrossPages:
        merged = []
        for t in tables:
            if merged and not t.get("noTable") and not merged[-1].get("noTable") \
                    and t["page"] == merged[-1]["lastPage"] + 1 and t["ncols"] == merged[-1]["ncols"]:
                prev = merged[-1]
                rows = list(t["rows"])
                trail = list(t.get("trail") or [[False] * len(r) for r in rows])
                if rows and prev["rows"] and rows[0] == prev["rows"][0]:
                    rows, trail = rows[1:], trail[1:]  # repeated header
                # A first row with an empty first column continues the previous row
                # (the generator split one row across the page break).
                if rows and prev["rows"] and not (rows[0][0] or "").strip() and any(c for c in rows[0][1:]):
                    last = prev["rows"][-1]
                    last_tr = prev.setdefault("trail", [[False] * len(r) for r in prev["rows"]])[-1]
                    for ci, frag in enumerate(rows[0]):
                        if not frag or ci >= len(last):
                            continue
                        joined, amb = join_wrapped(last[ci] or "", frag, last_tr[ci] if ci < len(last_tr) else False,
                                                   prev.get("keepsWrap", False))
                        last[ci] = joined
                        prev["stitched"] = prev.get("stitched", 0) + 1
                        if amb:
                            prev["ambiguous"] = prev.get("ambiguous", 0) + 1
                    rows, trail = rows[1:], trail[1:]
                prev["rows"].extend(rows)
                prev.setdefault("trail", []).extend(trail)
                prev["lastPage"] = t["page"]
                prev["pages"].append(t["page"])
                continue
            merged.append({**t, "lastPage": t["page"], "pages": [t["page"]]})
        tables = merged
    else:
        tables = [{**t, "lastPage": t["page"], "pages": [t["page"]]} for t in tables]

    if is_csv:
        return _write_csv(ctx, path, output, overwrite, tables, report, csvMode, ocr_info)

    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    converted = 0
    kept_text = 0

    def write_rows(ws, rows, start_row):
        nonlocal converted, kept_text
        r0 = start_row
        for ri, row in enumerate(rows):
            is_header = ri == 0 and len(rows) > 1 and any(
                c and parse_number(c) is None and parse_date(c, dateOrder) is None for c in row)
            for ci, val in enumerate(row):
                if val is None or val == "":
                    continue
                cell = ws.cell(row=r0 + ri, column=ci + 1)
                if convertNumbers and not is_header:
                    num = parse_number(val)
                    if num is not None:
                        cell.value, cell.number_format = num
                        converted += 1
                        continue
                    dt = parse_date(val, dateOrder)
                    if dt is not None:
                        cell.value, cell.number_format = dt
                        converted += 1
                        continue
                cell.value = val
                cell.number_format = "@"
                if digits_of(val) and len(digits_of(val)) >= 2:
                    kept_text += 1
        # Header styling and merged header cells (a value followed by empty cells).
        if rows and len(rows) > 1:
            for ci in range(len(rows[0])):
                c = ws.cell(row=r0, column=ci + 1)
                c.font = Font(bold=True)
                c.fill = PatternFill("solid", fgColor="E9ECEF")
                c.alignment = Alignment(wrap_text=True, vertical="top")
            hdr = rows[0]
            ci = 0
            while ci < len(hdr):
                if hdr[ci] and ci + 1 < len(hdr) and hdr[ci + 1] is None:
                    cj = ci + 1
                    while cj < len(hdr) and hdr[cj] is None:
                        cj += 1
                    ws.merge_cells(start_row=r0, start_column=ci + 1, end_row=r0, end_column=cj)
                    ci = cj
                else:
                    ci += 1
        return r0 + len(rows)

    def autosize(ws):
        widths = {}
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                L = len(str(c.value)) if not isinstance(c.value, (int, float)) else len(f"{c.value:,.2f}")
                widths[c.column] = max(widths.get(c.column, 0), min(60, L + 2))
        for col, w in widths.items():
            ws.column_dimensions[get_column_letter(col)].width = max(8, w)

    used_names = set()

    def sheet_name(base):
        base = re.sub(r"[\[\]:*?/\\]", "_", base)[:28]
        name = base
        k = 2
        while name in used_names:
            name = f"{base[:25]}_{k}"
            k += 1
        used_names.add(name)
        return name

    if layout == "singleSheet":
        ws = wb.create_sheet(sheet_name("Tables"))
        r = 1
        for ti, t in enumerate(tables):
            ws.cell(row=r, column=1, value=f"Table {ti + 1} (page {t['page']}" + (f"-{t['lastPage']}" if t['lastPage'] != t['page'] else "") + ")").font = Font(bold=True, italic=True)
            r = write_rows(ws, t["rows"], r + 1) + 1
            t["sheet"] = ws.title
        autosize(ws)
    elif layout == "sheetPerPage":
        by_page = {}
        for t in tables:
            by_page.setdefault(t["page"], []).append(t)
        for pno, ts in by_page.items():
            ws = wb.create_sheet(sheet_name(f"Page {pno}"))
            r = 1
            for t in ts:
                r = write_rows(ws, t["rows"], r) + 1
                t["sheet"] = ws.title
            autosize(ws)
    else:
        for ti, t in enumerate(tables):
            label = f"P{t['page']}" + (f"-{t['lastPage']}" if t["lastPage"] != t["page"] else "") + f" T{ti + 1}"
            if t.get("noTable"):
                label = f"P{t['page']} text"
            ws = wb.create_sheet(sheet_name(label))
            write_rows(ws, t["rows"], 1)
            ws.freeze_panes = "A2" if len(t["rows"]) > 1 and not t.get("noTable") else None
            autosize(ws)
            t["sheet"] = ws.title

    rs = wb.create_sheet("Extraction Report", 0)
    rs.append(["NayanPDF table extraction report"])
    rs["A1"].font = Font(bold=True, size=13)
    rs.append(["Source", os.path.basename(path)])
    rs.append(["Note", "Values were converted to numbers/dates only where the text matched an exact numeric or date pattern. "
                        "Identifiers (leading zeros, long references) are kept as text. Review flagged pages against the source."])
    rs.append([])
    rs.append(["Page", "Tables", "Detection", "Numbers on page", "Numbers captured", "Coverage %",
               "Wrapped cells to check", "Review", "Missing (sample)"])
    for c in rs[5]:
        c.font = Font(bold=True)
    for r in report:
        rs.append([r["page"], r["tables"], r["strategy"], r["numbersOnPage"], r["numbersCaptured"],
                   r["coverage"], r.get("ambiguousCells", 0), "REVIEW" if r["review"] else "ok",
                   ", ".join(r["missingSample"])])
        if r["review"]:
            rs.cell(row=rs.max_row, column=8).font = Font(bold=True, color="C92A2A")
    rs.append([])
    rs.append(["Sheets"])
    for t in tables:
        rs.append([t.get("sheet"), f"pages {', '.join(map(str, t['pages']))}", t["strategy"], f"{len(t['rows'])} rows",
                   f"{t['ncols']} columns"])
    autosize(rs)
    if other_text:
        ws = wb.create_sheet(sheet_name("Other text"))
        ws.append(["Page", "Text outside tables"])
        ws["A1"].font = Font(bold=True)
        ws["B1"].font = Font(bold=True)
        for pno, lines in other_text:
            for lt in lines:
                ws.append([pno, lt])
                ws.cell(row=ws.max_row, column=2).number_format = "@"
        ws.column_dimensions["B"].width = 100
    out = OutputFile(ctx, output, overwrite, sources=[path])
    wb.save(out.temp)
    final = out.commit(validate_xlsx)
    review = [r["page"] for r in report if r["review"]]
    if review:
        ctx.warn(f"Review recommended for page(s) {', '.join(map(str, review[:20]))} (see the 'Extraction Report' sheet).")
    res = {"output": final, "tables": len([t for t in tables if not t.get("noTable")]), "sheets": len(wb.sheetnames),
           "numbersConverted": converted, "numericTextKept": kept_text, "reviewPages": review, "report": report}
    if ocr_info:
        res["ocr"] = ocr_info
    return res


def _write_csv(ctx, path, output, overwrite, tables, report, mode, ocr_info):
    outputs = []
    real = [t for t in tables]
    if mode == "perTable" and len(real) > 1:
        base = os.path.splitext(output)[0]
        for k, t in enumerate(real):
            target = f"{base}_table{k + 1}.csv"
            out = OutputFile(ctx, target, overwrite, sources=[path])
            with open(out.temp, "w", encoding="utf-8-sig", newline="") as fh:
                w = csv.writer(fh)
                for r in t["rows"]:
                    w.writerow(["" if c is None else c for c in r])
            outputs.append(out.commit(validate_nonempty))
    else:
        out = OutputFile(ctx, output, overwrite, sources=[path])
        with open(out.temp, "w", encoding="utf-8-sig", newline="") as fh:
            w = csv.writer(fh)
            for k, t in enumerate(real):
                if len(real) > 1:
                    w.writerow([f"# Table {k + 1} (page {t['page']})"])
                for r in t["rows"]:
                    w.writerow(["" if c is None else c for c in r])
                if len(real) > 1:
                    w.writerow([])
        outputs.append(out.commit(validate_nonempty))
    review = [r["page"] for r in report if r["review"]]
    if review:
        ctx.warn(f"Review recommended for page(s) {', '.join(map(str, review[:20]))}.")
    return {"output": outputs[0], "outputs": outputs, "tables": len(real), "reviewPages": review, "report": report,
            "note": "CSV values are written exactly as extracted (no number conversion)."}


# ============================================================================
# PDF -> PowerPoint
# ============================================================================

EMU_PER_PT = 12700


def _font_family(pdf_font: str) -> str:
    name = pdf_font.split("+", 1)[-1]
    name = re.sub(r"[-,](Bold|Italic|Oblique|Regular|BoldItalic|BoldOblique|Roman|MT|PS|PSMT|Semibold|Light|Medium)+$", "", name, flags=re.I)
    name = re.sub(r"(PSMT|PS|MT)$", "", name)
    mapping = {"Helvetica": "Arial", "Times": "Times New Roman", "TimesNewRoman": "Times New Roman",
               "Courier": "Courier New", "Arial": "Arial", "Carlito": "Calibri", "Caladea": "Cambria"}
    return mapping.get(name, re.sub(r"(?<=[a-z])(?=[A-Z])", " ", name) or "Arial")


@method("convert.toPowerPoint")
def to_powerpoint(ctx, path: str, output: str, mode: str = "hybrid", pages=None, dpi: int = 200,
                  password: str | None = None, overwrite: bool = False):
    """mode: 'image' (page pictures, not editable), 'hybrid' (editable text over a
    background picture of the non-text content), 'editable' (text, images and
    simple shapes as individual objects)."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Emu, Pt

    if mode not in ("image", "hybrid", "editable"):
        raise invalid("mode must be image, hybrid or editable.")
    doc = open_pdf(path, password, require=["copy"])
    sel = parse_pages(pages, doc.page_count)
    prs = Presentation()
    first = doc[sel[0]].rect
    sw, sh = first.width, first.height
    # PowerPoint limits slides to 56 inches.
    scale_all = min(1.0, (56 * 72) / max(sw, sh))
    prs.slide_width = Emu(int(sw * scale_all * EMU_PER_PT))
    prs.slide_height = Emu(int(sh * scale_all * EMU_PER_PT))
    blank = prs.slide_layouts[6]
    stats = {"textBoxes": 0, "pictures": 0, "shapes": 0}
    for k, i in enumerate(sel):
        ctx.progress(k, len(sel), f"Slide {k + 1}")
        page = doc[i]
        pr = page.rect
        s = min(sw / pr.width, sh / pr.height) * scale_all
        ox = (sw * scale_all - pr.width * s) / 2
        oy = (sh * scale_all - pr.height * s) / 2

        def E(v):
            return Emu(int(v * EMU_PER_PT))

        slide = prs.slides.add_slide(blank)
        if mode == "image":
            pix = page.get_pixmap(dpi=int(dpi), alpha=False)
            slide.shapes.add_picture(io.BytesIO(pix.tobytes("png")), E(ox), E(oy), E(pr.width * s), E(pr.height * s))
            stats["pictures"] += 1
            continue
        # Text lines in displayed orientation.
        lines = []
        rm = page.rotation_matrix
        for b in page.get_text("dict")["blocks"]:
            for ln in b.get("lines", []):
                txt = "".join(sp["text"] for sp in ln["spans"])
                if not txt.strip():
                    continue
                bb = pymupdf.Rect(ln["bbox"]) * rm
                lines.append((bb, ln))
        if mode == "hybrid":
            bg = _render_without_text(page, int(dpi))
            slide.shapes.add_picture(io.BytesIO(bg), E(ox), E(oy), E(pr.width * s), E(pr.height * s))
            stats["pictures"] += 1
        else:
            drawings = page.get_drawings()
            if len(drawings) > 600:
                bg = _render_without_text(page, int(dpi), remove_images=True)
                slide.shapes.add_picture(io.BytesIO(bg), E(ox), E(oy), E(pr.width * s), E(pr.height * s))
                stats["pictures"] += 1
                ctx.warn(f"Page {i + 1} has complex vector graphics; they were placed as a background picture.")
            else:
                from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE

                for dr in drawings:
                    r = pymupdf.Rect(dr["rect"]) * rm
                    if r.is_empty and r.width < 0.5 and r.height < 0.5:
                        continue
                    fill = dr.get("fill")
                    color = dr.get("color")
                    items = dr.get("items", [])
                    if len(items) == 1 and items[0][0] == "l":
                        p1 = pymupdf.Point(items[0][1]) * rm
                        p2 = pymupdf.Point(items[0][2]) * rm
                        cn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(ox + p1.x * s), E(oy + p1.y * s),
                                                        E(ox + p2.x * s), E(oy + p2.y * s))
                        if color:
                            cn.line.color.rgb = RGBColor.from_string(rgb_to_hex(color)[1:])
                        cn.line.width = Pt(max(0.25, (dr.get("width") or 1) * s))
                        stats["shapes"] += 1
                    elif all(it[0] == "re" for it in items) or (len(items) == 4 and all(it[0] == "l" for it in items)):
                        if r.width < 0.5 or r.height < 0.5:
                            continue
                        shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, E(ox + r.x0 * s), E(oy + r.y0 * s),
                                                     E(max(0.5, r.width) * s), E(max(0.5, r.height) * s))
                        if fill:
                            shp.fill.solid()
                            shp.fill.fore_color.rgb = RGBColor.from_string(rgb_to_hex(fill)[1:])
                        else:
                            shp.fill.background()
                        if color and dr.get("width"):
                            shp.line.color.rgb = RGBColor.from_string(rgb_to_hex(color)[1:])
                            shp.line.width = Pt(max(0.25, dr["width"] * s))
                        else:
                            shp.line.fill.background()
                        shp.shadow.inherit = False
                        stats["shapes"] += 1
            for img in page.get_image_info(xrefs=True):
                xref = img.get("xref")
                r = pymupdf.Rect(img["bbox"]) * rm
                if r.is_empty or r.width < 2 or r.height < 2:
                    continue
                try:
                    if xref:
                        pix = pymupdf.Pixmap(doc, xref)
                        if pix.colorspace and pix.colorspace.n > 3:
                            pix = pymupdf.Pixmap(pymupdf.csRGB, pix)
                        data = pix.tobytes("png")
                    else:
                        data = page.get_pixmap(clip=pymupdf.Rect(img["bbox"]), dpi=int(dpi)).tobytes("png")
                    slide.shapes.add_picture(io.BytesIO(data), E(ox + r.x0 * s), E(oy + r.y0 * s), E(r.width * s), E(r.height * s))
                    stats["pictures"] += 1
                except Exception:  # noqa: BLE001
                    continue
        for bb, ln in lines:
            if abs(ln["dir"][1]) > 0.01:
                continue
            tb = slide.shapes.add_textbox(E(ox + bb.x0 * s), E(oy + bb.y0 * s), E(max(4, bb.width * s * 1.04)), E(max(4, bb.height * s)))
            tf = tb.text_frame
            tf.word_wrap = False
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            p = tf.paragraphs[0]
            for sp in ln["spans"]:
                if not sp["text"]:
                    continue
                run = p.add_run()
                run.text = sp["text"]
                run.font.size = Pt(max(1, round(sp["size"] * s * 2) / 2))
                run.font.bold = bool(sp["flags"] & 16) or "bold" in sp["font"].lower()
                run.font.italic = bool(sp["flags"] & 2) or "italic" in sp["font"].lower()
                run.font.name = _font_family(sp["font"])
                run.font.color.rgb = RGBColor.from_string(rgb_to_hex(int_to_rgb(sp["color"]))[1:])
            stats["textBoxes"] += 1
    doc.close()
    out = OutputFile(ctx, output, overwrite, sources=[path])
    prs.save(out.temp)
    final = out.commit(validate_pptx)
    label = {"image": "Slides contain page images (not editable).",
             "hybrid": "Text is editable; graphics and images are a background picture on each slide.",
             "editable": "Text, images and simple shapes are separate editable objects; complex graphics may differ."}[mode]
    if mode != "image" and stats["textBoxes"] == 0:
        ctx.warn("No text was found to make editable (scanned pages need OCR first).")
    return {"output": final, "slides": len(sel), "mode": mode, "description": label, **stats}


def _render_without_text(page: pymupdf.Page, dpi: int, remove_images: bool = False) -> bytes:
    tmp = pymupdf.open()
    tmp.insert_pdf(page.parent, from_page=page.number, to_page=page.number, annots=True)
    p = tmp[0]
    unrot = p.rect * p.derotation_matrix
    unrot.normalize()
    p.add_redact_annot(unrot, fill=False, cross_out=False)
    p.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_REMOVE if remove_images else pymupdf.PDF_REDACT_IMAGE_NONE,
                       graphics=pymupdf.PDF_REDACT_LINE_ART_NONE, text=pymupdf.PDF_REDACT_TEXT_REMOVE)
    pix = p.get_pixmap(dpi=dpi, alpha=False)
    data = pix.tobytes("png")
    tmp.close()
    return data


# ============================================================================
# PDF -> images
# ============================================================================

@method("convert.toImages")
def to_images(ctx, path: str, outputDir: str, format: str = "png", dpi: int = 150, quality: int = 90,
              pages=None, nameTemplate: str = "{name}_page_{page:03d}", grayscale: bool = False,
              multiPageTiff: bool = False, password: str | None = None, overwrite: bool = False):
    fmt = format.lower()
    if fmt == "jpeg":
        fmt = "jpg"
    if fmt not in ("png", "jpg", "tiff", "bmp", "webp"):
        raise invalid("Format must be PNG, JPG, TIFF, BMP or WebP.")
    dpi = max(36, min(1200, int(dpi)))
    doc = open_pdf(path, password, require=["copy"])
    sel = parse_pages(pages, doc.page_count)
    outputs = []
    from PIL import Image

    pil_frames = []
    for k, i in enumerate(sel):
        ctx.progress(k, len(sel), f"Page {i + 1}")
        page = doc[i]
        pix = page.get_pixmap(dpi=dpi, colorspace=pymupdf.csGRAY if grayscale else pymupdf.csRGB, alpha=False)
        exp = (pix.width, pix.height)
        if multiPageTiff and fmt == "tiff":
            pil_frames.append(Image.frombytes("L" if grayscale else "RGB", exp, pix.samples))
            continue
        name = apply_name_template(nameTemplate, path, page=i + 1, total=doc.page_count, n=k + 1)
        target = os.path.join(outputDir, f"{name}.{fmt}")
        out = OutputFile(ctx, target, overwrite, sources=[path])
        if fmt == "png":
            pix.save(out.temp)
        elif fmt == "jpg":
            pix.save(out.temp, jpg_quality=max(10, min(100, int(quality))))
        else:
            im = Image.frombytes("L" if grayscale else "RGB", exp, pix.samples)
            if fmt == "tiff":
                im.save(out.temp, "TIFF", compression="tiff_lzw", dpi=(dpi, dpi))
            elif fmt == "bmp":
                im.save(out.temp, "BMP")
            else:
                im.save(out.temp, "WEBP", quality=max(10, min(100, int(quality))), method=4)

        def check(p, exp=exp):
            info = validate_image(p)
            if (info["width"], info["height"]) != exp:
                raise EngineError(VALIDATION_FAILED, "Exported image has unexpected dimensions.")
            return info

        final = out.commit(check)
        outputs.append({"output": final, "page": i + 1, "width": exp[0], "height": exp[1]})
    if multiPageTiff and fmt == "tiff":
        target = derive_output(path, outputDir, "", ".tiff")
        out = OutputFile(ctx, target, overwrite, sources=[path])
        pil_frames[0].save(out.temp, "TIFF", compression="tiff_lzw", save_all=True, append_images=pil_frames[1:],
                           dpi=(dpi, dpi))
        final = out.commit(validate_image)
        outputs.append({"output": final, "pages": len(pil_frames)})
    doc.close()
    return {"outputs": outputs, "count": len(outputs), "format": fmt, "dpi": dpi}


# ============================================================================
# PDF -> text / HTML / Markdown
# ============================================================================

def _layout_text(page: pymupdf.Page) -> str:
    """Approximate the visual layout with spaces (keeps table columns aligned)."""
    words = page.get_text("words", sort=True)
    if not words:
        return ""
    widths = [(w[2] - w[0]) / max(1, len(w[4])) for w in words if w[4]]
    cw = max(2.0, statistics.median(widths)) if widths else 5.0
    lines: list[list] = []
    for w in sorted(words, key=lambda w: ((w[1] + w[3]) / 2, w[0])):
        yc = (w[1] + w[3]) / 2
        if lines and abs(lines[-1][0] - yc) < (w[3] - w[1]) * 0.5:
            lines[-1][1].append(w)
        else:
            lines.append([yc, [w]])
    out = []
    x_min = min(w[0] for w in words)
    for _, ws in lines:
        ws.sort(key=lambda w: w[0])
        line = ""
        for w in ws:
            col = int(round((w[0] - x_min) / cw))
            if len(line) < col:
                line += " " * (col - len(line))
            elif line and not line.endswith(" "):
                line += " "
            line += w[4]
        out.append(line.rstrip())
    return "\n".join(out)


@method("convert.toText")
def to_text(ctx, path: str, output: str, pages=None, layout: bool = True, ocr: str = "auto",
            languages: list | None = None, password: str | None = None, overwrite: bool = False):
    src, pw, ocr_info = _maybe_ocr(ctx, path, password, ocr, languages)
    doc = open_pdf(src, pw, require=["copy"])
    sel = parse_pages(pages, doc.page_count)
    out = OutputFile(ctx, output, overwrite, sources=[path])
    chars = 0
    with open(out.temp, "w", encoding="utf-8", newline="\r\n") as fh:
        for k, i in enumerate(sel):
            ctx.progress(k, len(sel))
            t = _layout_text(doc[i]) if layout else doc[i].get_text("text", sort=True)
            chars += len(t.strip())
            fh.write(t)
            fh.write("\n\f" if k < len(sel) - 1 else "\n")
    doc.close()
    final = out.commit(validate_nonempty)
    if chars == 0:
        ctx.warn("No text was extracted. The PDF may be scanned; enable OCR.")
    res = {"output": final, "characters": chars}
    if ocr_info:
        res["ocr"] = ocr_info
    return res


@method("convert.toHtml")
def to_html(ctx, path: str, output: str, mode: str = "semantic", pages=None, password: str | None = None,
            overwrite: bool = False):
    import html as _h

    doc = open_pdf(path, password, require=["copy"])
    sel = parse_pages(pages, doc.page_count)
    title = (doc.metadata or {}).get("title") or os.path.splitext(os.path.basename(path))[0]
    parts = [f"<!DOCTYPE html>\n<html><head><meta charset=\"utf-8\"><title>{_h.escape(title)}</title>",
             "<style>body{font-family:Segoe UI,Arial,sans-serif;max-width:960px;margin:24px auto;padding:0 16px;color:#222}"
             ".page{margin:24px 0;border-bottom:1px solid #ddd;padding-bottom:16px;position:relative}"
             "table{border-collapse:collapse}td,th{border:1px solid #bbb;padding:2px 6px}</style></head><body>"]
    for k, i in enumerate(sel):
        ctx.progress(k, len(sel))
        page = doc[i]
        if mode == "layout":
            parts.append(f'<section class="page" id="page{i + 1}">' + page.get_text("html") + "</section>")
        else:
            parts.append(f'<section class="page" id="page{i + 1}">' + page.get_text("xhtml") + "</section>")
    parts.append("</body></html>")
    doc.close()
    out = OutputFile(ctx, output, overwrite, sources=[path])
    with open(out.temp, "w", encoding="utf-8") as fh:
        fh.write("\n".join(parts))
    final = out.commit(validate_nonempty)
    return {"output": final, "mode": mode}


@method("convert.toMarkdown")
def to_markdown(ctx, path: str, output: str, pages=None, ocr: str = "auto", languages: list | None = None,
                password: str | None = None, overwrite: bool = False):
    src, pw, ocr_info = _maybe_ocr(ctx, path, password, ocr, languages)
    doc = open_pdf(src, pw, require=["copy"])
    sel = parse_pages(pages, doc.page_count)
    sizes = []
    for i in sel[:30]:
        for b in doc[i].get_text("dict")["blocks"]:
            for ln in b.get("lines", []):
                for s in ln["spans"]:
                    if s["text"].strip():
                        sizes.append(round(s["size"]))
    body = statistics.mode(sizes) if sizes else 11
    md = []
    for k, i in enumerate(sel):
        ctx.progress(k, len(sel))
        page = doc[i]
        try:
            tabs = [t for t in page.find_tables().tables if t.row_count >= 2 and t.col_count >= 2]
        except Exception:  # noqa: BLE001
            tabs = []
        tboxes = [pymupdf.Rect(t.bbox) for t in tabs]
        items = [(t.bbox[1], "t", t) for t in tabs]
        for b in page.get_text("dict", sort=True)["blocks"]:
            if b["type"] != 0:
                continue
            bb = pymupdf.Rect(b["bbox"])
            if any((bb & tb).get_area() > 0.5 * bb.get_area() for tb in tboxes):
                continue
            items.append((bb.y0, "b", b))
        items.sort(key=lambda x: x[0])
        for _, kind, obj in items:
            if kind == "t":
                try:
                    md.append(obj.to_markdown(clean=False).strip())
                except Exception:  # noqa: BLE001
                    rows = obj.extract()
                    md.append("\n".join("| " + " | ".join((c or "").replace("\n", " ") for c in r) + " |" for r in rows))
                md.append("")
                continue
            spans = [s for ln in obj["lines"] for s in ln["spans"] if s["text"].strip()]
            if not spans:
                continue
            text = " ".join("".join(s["text"] for s in ln["spans"]).strip() for ln in obj["lines"]).strip()
            size = max(s["size"] for s in spans)
            bold = all((s["flags"] & 16) for s in spans)
            if size >= body * 1.45 and len(text) < 200:
                md.append(f"# {text}\n")
            elif size >= body * 1.2 and len(text) < 200:
                md.append(f"## {text}\n")
            elif bold and len(text) < 120:
                md.append(f"### {text}\n")
            elif re.match(r"^([•●\-\*]|\d+[.)])\s+", text):
                md.append(re.sub(r"^[•●]\s*", "- ", text))
            else:
                md.append(text + "\n")
        if k < len(sel) - 1:
            md.append("\n---\n")
    doc.close()
    out = OutputFile(ctx, output, overwrite, sources=[path])
    with open(out.temp, "w", encoding="utf-8") as fh:
        fh.write("\n".join(md).strip() + "\n")
    final = out.commit(validate_nonempty)
    res = {"output": final}
    if ocr_info:
        res["ocr"] = ocr_info
    return res


# ============================================================================
# Office -> PDF (LibreOffice, headless, private profile)
# ============================================================================

OFFICE_EXT = {".doc", ".docx", ".odt", ".rtf", ".xls", ".xlsx", ".ods", ".csv", ".ppt", ".pptx", ".odp",
              ".txt", ".wpd", ".docm", ".xlsm", ".pptm", ".dot", ".dotx", ".xlt", ".xltx", ".pot", ".potx"}


def office_to_pdf_file(ctx, src: str, target_tmp: str, singlePageSheets: bool = False) -> None:
    exe = tools.require_tool(tools.soffice_exe(), "LibreOffice (Office conversion)")
    ext = os.path.splitext(src)[1].lower()
    if ext not in OFFICE_EXT:
        raise EngineError(UNSUPPORTED_FORMAT, f"{ext or 'This file type'} is not a supported Office document.")
    work = os.path.join(ctx.tmpdir, "lo")
    os.makedirs(work, exist_ok=True)
    profile, release = _lo_profile(ctx)
    # Copy the source to an ASCII-safe name (LibreOffice is fragile with some paths).
    tmp_src = os.path.join(work, "input" + ext)
    shutil.copyfile(os_path(src), tmp_src)
    filt = "pdf"
    if ext in (".xls", ".xlsx", ".ods", ".csv", ".xlsm", ".xlt", ".xltx") and singlePageSheets:
        filt = 'pdf:calc_pdf_Export:{"SinglePageSheets":{"type":"boolean","value":"true"}}'
    import pathlib

    profile_url = pathlib.Path(profile).resolve().as_uri()
    args = [exe, f"-env:UserInstallation={profile_url}", "--headless", "--norestore", "--nologo",
            "--nodefault", "--nolockcheck", "--convert-to", filt, "--outdir", work, tmp_src]
    try:
        res = tools.run(args, timeout=900, env={"SAL_USE_VCLPLUGIN": "svp"} if os.name != "nt" else None)
    finally:
        release()
    produced = os.path.join(work, "input.pdf")
    if not os.path.exists(produced) or os.path.getsize(produced) == 0:
        msg = (res.stderr or b"").decode("utf-8", "replace")[-300:]
        raise EngineError("OPERATION_FAILED", f"LibreOffice could not convert {os.path.basename(src)}. {msg}".strip())
    shutil.move(produced, target_tmp)


def _lo_profile(ctx):
    """A reusable LibreOffice user profile (creating one takes several seconds).

    Profiles live under NAYANPDF_LO_PROFILES (set by the app to its data folder);
    a lock file per slot keeps concurrent conversions from sharing a profile."""
    import time

    base = os.environ.get("NAYANPDF_LO_PROFILES")
    if not base:
        p = os.path.join(ctx.tmpdir, "lo-profile")
        os.makedirs(p, exist_ok=True)
        return p, lambda: None
    os.makedirs(base, exist_ok=True)
    for slot in range(8):
        lock = os.path.join(base, f"slot{slot}.lock")
        try:
            if os.path.exists(lock) and time.time() - os.path.getmtime(lock) > 3600:
                os.remove(lock)  # stale lock from a crashed conversion
            fd = os.open(lock, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
            os.close(fd)
        except OSError:
            continue
        prof = os.path.join(base, f"slot{slot}")
        os.makedirs(prof, exist_ok=True)

        def release(lock=lock):
            try:
                os.remove(lock)
            except OSError:
                pass

        return prof, release
    p = os.path.join(ctx.tmpdir, "lo-profile")
    os.makedirs(p, exist_ok=True)
    return p, lambda: None


@method("convert.officeToPdf")
def office_to_pdf(ctx, path: str, output: str, singlePageSheets: bool = False, overwrite: bool = False):
    require_file(path)
    out = OutputFile(ctx, output, overwrite, sources=[path])
    ctx.status("Converting with LibreOffice")
    office_to_pdf_file(ctx, path, out.temp, singlePageSheets)
    final = out.commit(validate_pdf())
    d = pymupdf.open(final)
    n = d.page_count
    d.close()
    ctx.warn("Office documents are rendered by LibreOffice. Fonts not installed on this computer are substituted, which can change line breaks; compare the PDF with the original.")
    return {"output": final, "pageCount": n}
